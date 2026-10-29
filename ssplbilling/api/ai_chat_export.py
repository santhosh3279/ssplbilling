"""Excel files and charts for the AI Chat, built from the data tools in ai_chat_tools.

The model (or a quick button) only names the source — a data tool and its args. The data is
re-read here, as the chat user, and goes straight into the file or chart, so a 5,000-row
export costs the same tokens as a 10-row answer and no figure is ever typed by the model.
Results are appended to `attachments` as dicts the chat page renders:

	{"type": "excel", "file": <File name>, "file_name": ..., "title": ..., "rows": n}
	{"type": "chart", "chart_type": "bar"|"line"|"pie", "title": ..., "labels": [...],
	 "datasets": [{"name": ..., "values": [...]}]}
"""

import io
import re

import frappe
from frappe.utils import cint, flt, getdate, random_string
from openpyxl import Workbook
from openpyxl.styles import Font

from ssplbilling.api import ai_chat_tools as tools

FILE_PREFIX = "ai-chat-"
BAR_POINTS = 30
LINE_POINTS = 400

LABELS = {
	"company": "Company",
	"customer": "Customer",
	"customer_name": "Customer Name",
	"mobile_no": "Mobile",
	"item_code": "Item Code",
	"item_name": "Item Name",
	"stock_uom": "UOM",
	"uom": "UOM",
	"warehouse": "Warehouse",
	"posting_date": "Date",
	"due_date": "Due Date",
	"month": "Month",
	"naming_series": "Series",
	"name": "Invoice",
	"invoices": "Invoices",
	"qty": "Qty",
	"net_amount": "Net Amount (₹)",
	"grand_total": "Grand Total (₹)",
	"net_total": "Net Total (₹)",
	"balance": "Balance (₹)",
	"outstanding_amount": "Outstanding (₹)",
	"days_overdue": "Days Overdue",
	"items_in_stock": "Items in Stock",
	"items_negative": "Items Negative",
	"total_qty": "Qty",
	"actual_qty": "Qty",
	"reserved_qty": "Reserved",
	"projected_qty": "Projected",
	"valuation_rate": "Valuation Rate (₹)",
	"stock_value": "Stock Value (₹)",
	"receivable": "Receivable (₹)",
	"advances": "Advances (₹)",
	"customers_owing": "Customers Owing",
	"returns": "Returns (₹)",
	"net_sales": "Net Sales (₹)",
	"draft_invoices": "Draft Invoices",
	"draft_total": "Draft Total (₹)",
	"account": "Account",
	"voucher_type": "Voucher Type",
	"voucher_no": "Voucher",
	"party": "Party",
	"against": "Against",
	"remarks": "Remarks",
	"day": "Date",
	"vouchers": "Vouchers",
	"debit": "Debit (₹)",
	"credit": "Credit (₹)",
	"net": "Net (₹)",
	"amount": "Amount (₹)",
	"opening_balance": "Opening (₹)",
	"closing_balance": "Closing (₹)",
	"against_total": "Against Total (₹)",
	"against_amount": "Against Amount (₹)",
	"supplier": "Supplier",
	"supplier_name": "Supplier Name",
	"bill_no": "Bill No",
	"payable": "Payable (₹)",
	"suppliers_owed": "Suppliers Owed",
	"net_purchases": "Net Purchases (₹)",
	"mode_of_payment": "Mode",
	"party_type": "Party Type",
	"party_name": "Party Name",
	"entries": "Entries",
	"group": "Group",
	"type": "Type",
	"total": "Total (₹)",
	"income": "Income (₹)",
	"expense": "Expense (₹)",
	"net_profit": "Net Profit (₹)",
	"net_payable": "Net Payable (₹)",
	"qty_sold": "Qty Sold",
	"in_qty": "In",
	"out_qty": "Out",
	"opening": "Opening Qty",
	"closing": "Closing Qty",
	"last_sold": "Last Sold",
	"last_purchase": "Last Purchase",
	"first_invoice": "First Invoice",
	"last_invoice": "Last Invoice",
	"days_since": "Days Since",
	"days_since_sale": "Days Since Sale",
	"days_since_last": "Days Since Last",
}
MONEY = {
	"net_amount",
	"grand_total",
	"net_total",
	"balance",
	"outstanding_amount",
	"valuation_rate",
	"stock_value",
	"receivable",
	"advances",
	"returns",
	"net_sales",
	"draft_total",
	"debit",
	"credit",
	"net",
	"amount",
	"opening_balance",
	"closing_balance",
	"against_total",
	"against_amount",
	"payable",
	"net_purchases",
	"total",
	"income",
	"expense",
	"net_profit",
	"net_payable",
}
# Generic sheets flatten nested totals into columns like received_total or output_by_tax_CGST.
MONEY_SUFFIXES = ("total", "amount", "balance", "value", "output", "input", "CGST", "SGST", "IGST", "CESS")
QTY = {
	"qty",
	"total_qty",
	"actual_qty",
	"reserved_qty",
	"projected_qty",
	"qty_sold",
	"in_qty",
	"out_qty",
	"opening",
	"closing",
}
DATES = {
	"posting_date",
	"due_date",
	"last_sold",
	"last_purchase",
	"first_invoice",
	"last_invoice",
	"as_of",
	"from_date",
	"to_date",
}
# Explanations meant for the model, not columns.
SKIP_KEYS = {"note", "sign", "amounts", "found", "against_note", "valuation_shown"}
LEDGER_SHEETS = {
	"by_against": "Against Accounts",
	"entries": "Entries",
	"by_day": "By Day",
	"by_month": "By Month",
}
INVOICE_TOOLS = {"sales_summary": "Sales Invoice", "purchase_summary": "Purchase Invoice"}
PARTY_TOOLS = {"customer_outstanding": "Customer", "supplier_outstanding": "Supplier"}
FORMATS = {"money": "#,##0.00", "qty": "#,##0.###", "date": "DD-MM-YYYY"}


# ── tools the model calls ───────────────────────────────────────────────────


def export_excel(tool, title, args=None, attachments=None):
	sheets = build_sheets(tool, args)
	rows = sum(len(s["rows"]) for s in sheets)
	if not rows:
		return {"ok": False, "message": "No rows to export."}
	file = _save_xlsx(sheets, title)
	if attachments is not None:
		attachments.append(
			{"type": "excel", "file": file.name, "file_name": file.file_name, "title": title, "rows": rows}
		)
	return {"ok": True, "file_name": file.file_name, "rows": rows, "sheets": [s["name"] for s in sheets]}


def make_chart(tool, title, chart_type="bar", args=None, attachments=None):
	if chart_type not in ("bar", "line", "pie"):
		raise ValueError(f"chart_type must be bar, line or pie; got {chart_type!r}")
	labels, datasets = build_series(tool, args, chart_type)
	if not labels:
		return {"ok": False, "message": "No data to chart."}
	if attachments is not None:
		attachments.append(
			{
				"type": "chart",
				"chart_type": chart_type,
				"title": title,
				"labels": labels,
				"datasets": datasets,
			}
		)
	return {"ok": True, "points": len(labels), "series": [d["name"] for d in datasets]}


# ── tables ──────────────────────────────────────────────────────────────────


def _run(tool, args, limit):
	if tool not in tools.DATA_TOOLS:
		raise ValueError(f"tool must be one of {', '.join(tools.DATA_TOOLS)}; got {tool!r}")
	params = _params(tool)
	args = {k: v for k, v in (args or {}).items() if k in params}
	if "limit" in params:
		args["limit"] = limit
	with tools.export_rows():
		data = tools.HANDLERS[tool](**args)
	if data.get("found") is False:
		raise ValueError(data.get("message") or "Nothing found.")
	return data


def _params(tool):
	return tools.HANDLERS[tool].__code__.co_varnames[: tools.HANDLERS[tool].__code__.co_argcount]


def build_sheets(tool, args=None):
	"""The tool's data as [{"name", "rows": [dict]}], one sheet per company where it splits."""
	data = _run(tool, args, tools.EXPORT_ROWS)
	if tool in INVOICE_TOOLS:
		cfg = tools.INVOICES[INVOICE_TOOLS[tool]]
		summary = [
			{
				"company": company,
				"invoices": c[cfg["total"]]["invoices"],
				"grand_total": c[cfg["total"]]["grand_total"],
				"returns": c["returns"]["grand_total"],
				cfg["net"]: c[cfg["net"]],
				"draft_invoices": c["pending_drafts"]["invoices"],
				"draft_total": c["pending_drafts"]["grand_total"],
			}
			for company, c in data["companies"].items()
		]
		sheets = [{"name": "Summary", "rows": summary}]
		for company, c in data["companies"].items():
			for key, rows in c.items():
				if key.startswith("by_"):
					sheets.append({"name": company, "rows": rows})
		return sheets

	if tool in PARTY_TOOLS:
		cfg = tools.PARTIES[PARTY_TOOLS[tool]]
		if data.get(cfg["key"]):
			return [
				{
					"name": "Balance",
					"rows": [{"company": co, "balance": b} for co, b in data["balance_by_company"].items()],
				},
				{"name": "Open Invoices", "rows": data["open_invoices"]},
			]
		summary = [
			{
				"company": company,
				cfg["total"]: c[cfg["total"]],
				"advances": c["advances"],
				cfg["count"]: c[cfg["count"]],
			}
			for company, c in data["companies"].items()
		]
		return [{"name": "Summary", "rows": summary}] + [
			{"name": company, "rows": c["top"]} for company, c in data["companies"].items()
		]

	if tool == "account_ledger":
		keys = (
			"account",
			"company",
			"opening_balance",
			"debit",
			"credit",
			"closing_balance",
			"against_total",
		)
		summary = {k: data[k] for k in keys if data.get(k) is not None}
		sheets = [{"name": "Summary", "rows": [summary]}]
		for key, name in LEDGER_SHEETS.items():
			if data.get(key):
				sheets.append({"name": name, "rows": data[key]})
		return sheets

	if tool == "item_stock":
		if "warehouses" in data:
			return [{"name": "Warehouses", "rows": data["warehouses"]}]
		rows = []
		for item in data["items"]:
			for w in item["warehouses"]:
				rows.append(
					{"item_code": item["item_code"], "item_name": item["item_name"], "uom": item["uom"], **w}
				)
		return [{"name": "Stock", "rows": rows}]

	return _generic_sheets(data)


def _generic_sheets(data):
	"""Any tool's result as a Summary sheet of its totals plus one sheet per list of rows."""
	companies = data.get("companies")
	if not isinstance(companies, dict):
		row, lists = _split(data)
		return [{"name": "Summary", "rows": [row]}, *lists]
	summary, sheets = [], []
	for company, c in companies.items():
		row, lists = _split(c)
		summary.append({"company": company, **row})
		abbr = frappe.get_cached_value("Company", company, "abbr") if len(companies) > 1 else None
		sheets += [{**sheet, "name": f"{abbr} {sheet['name']}" if abbr else sheet["name"]} for sheet in lists]
	return [{"name": "Summary", "rows": summary}, *sheets]


def _split(data, prefix=""):
	"""(scalar fields, nested dicts flattened to prefix_key; [{"name", "rows"}] for each list of rows)."""
	row, lists = {}, []
	for key, value in data.items():
		if key in SKIP_KEYS:
			continue
		if isinstance(value, list):
			if value and isinstance(value[0], dict):
				lists.append({"name": f"{prefix}{key}".replace("_", " ").title(), "rows": value})
		elif isinstance(value, dict):
			nested_row, nested_lists = _split(value, f"{prefix}{key}_")
			row.update(nested_row)
			lists += nested_lists
		else:
			row[f"{prefix}{key}"] = value
	return row, lists


def _save_xlsx(sheets, title):
	wb = Workbook()
	wb.remove(wb.active)
	used = set()
	for sheet in sheets:
		if not sheet["rows"]:
			continue
		ws = wb.create_sheet(_sheet_title(sheet["name"], used))
		keys = _columns(sheet["rows"])
		ws.append([LABELS.get(k, k.replace("_", " ").title()) for k in keys])
		for cell in ws[1]:
			cell.font = Font(bold=True)
		for row in sheet["rows"]:
			# Customers named by their own code have no customer_name row value: repeat the code.
			if "customer_name" in keys and not row.get("customer_name"):
				row = {**row, "customer_name": row.get("customer")}
			ws.append([_cell(k, row.get(k)) for k in keys])
		for i, key in enumerate(keys, start=1):
			kind = _kind(key)
			if kind:
				for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
					cell.number_format = FORMATS[kind]
			width = max(
				len(str(ws.cell(row=r, column=i).value or "")) for r in range(1, min(ws.max_row, 200) + 1)
			)
			ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(max(width + 2, 8), 50)
		ws.freeze_panes = "A2"
	wb.properties.title = title

	buf = io.BytesIO()
	wb.save(buf)
	return frappe.get_doc(
		{
			"doctype": "File",
			"file_name": f"{FILE_PREFIX}{re.sub(r'[^a-z0-9]+', '_', title.lower()).strip('_')[:40] or 'export'}-{random_string(6)}.xlsx",
			"is_private": 1,
			"content": buf.getvalue(),
		}
	).insert(ignore_permissions=True)


def _kind(key):
	if key in MONEY or key.rsplit("_", 1)[-1] in MONEY_SUFFIXES:
		return "money"
	if key in QTY:
		return "qty"
	if key in DATES:
		return "date"
	return None


def _columns(rows):
	keys = []
	for row in rows:
		keys += [k for k in row if k not in keys]
	return keys


def _cell(key, value):
	if key in DATES and value:
		return getdate(value)
	if isinstance(value, bool):
		return "Yes" if value else "No"
	if value is None or isinstance(value, int | float | str):
		return value
	return str(value)


def _sheet_title(name, used):
	title = re.sub(r"[\[\]:*?/\\]", " ", str(name))[:31].strip() or "Sheet"
	base, n = title, 2
	while title in used:
		title = f"{base[:28]} {n}"
		n += 1
	used.add(title)
	return title


# ── charts ──────────────────────────────────────────────────────────────────


def build_series(tool, args=None, chart_type="bar"):
	"""(labels, [{"name", "values"}]) for a chart, one dataset per company where it splits."""
	args = dict(args or {})
	points = LINE_POINTS if chart_type == "line" else BAR_POINTS
	# A line (trend) needs every point; a bar or pie shows the top N the user asked for.
	limit = points if chart_type == "line" else min(cint(args.get("limit")) or 10, points)
	data = _run(tool, args, limit)

	if tool in INVOICE_TOOLS:
		cfg = tools.INVOICES[INVOICE_TOOLS[tool]]
		total, net = cfg["total"], cfg["net"]
		group = args.get("group_by") or "none"
		if group == "none":
			labels = list(data["companies"])
			if chart_type == "pie":
				return labels, [
					{"name": f"Net {total}", "values": [c[net] for c in data["companies"].values()]}
				]
			return labels, [
				{
					"name": name,
					"values": [
						c[key] if key == net else c[key]["grand_total"] for c in data["companies"].values()
					],
				}
				for name, key in ((total.title(), total), ("Returns", "returns"), ("Net", net))
			]
		label_key = {
			"customer": "customer_name",
			"supplier": "supplier_name",
			"item": "item_name",
			"day": "posting_date",
			"month": "month",
			"series": "naming_series",
		}[group]
		value_key = "net_amount" if group == "item" else "grand_total"
		per_company = {
			company: {_label(r, label_key): r[value_key] for r in c.get(f"by_{group}", [])}
			for company, c in data["companies"].items()
		}
		return _align(per_company, points, sort=group in ("day", "month"))

	if tool == "top_items":
		return _align(
			{
				co: {_label(r, "item_name"): r["net_amount"] for r in c["items"]}
				for co, c in data["companies"].items()
			},
			points,
		)

	if tool in PARTY_TOOLS:
		key = tools.PARTIES[PARTY_TOOLS[tool]]["key"]
		if data.get(key):
			balances = data["balance_by_company"]
			return list(balances), [{"name": "Balance", "values": list(balances.values())}]
		per_company = {
			company: {_label(r, f"{key}_name"): r["balance"] for r in c["top"]}
			for company, c in data["companies"].items()
		}
		return _align(per_company, points)

	if tool == "account_ledger":
		if data.get("by_against") is not None:
			rows = data["by_against"]
			return [r["account"] for r in rows], [{"name": "Amount", "values": [r["amount"] for r in rows]}]
		if data.get("entries") is not None:
			rows = data["entries"]
			key, name = ("against_amount", "Amount") if data.get("against_filter") else ("balance", "Balance")
			return [r["voucher_no"] for r in rows], [{"name": name, "values": [r.get(key, 0) for r in rows]}]
		period = "day" if "by_day" in data else "month"
		rows = data[f"by_{period}"]
		key, name = ("amount", data["against_filter"]) if data.get("against_filter") else ("net", "Net")
		return [r[period] for r in rows], [{"name": name, "values": [r.get(key, 0) for r in rows]}]

	if tool == "item_stock":
		if "warehouses" in data:
			rows = data["warehouses"][:points]
			key, name = ("stock_value", "Stock value") if data["valuation_shown"] else ("total_qty", "Qty")
			return [r["warehouse"] for r in rows], [
				{"name": name, "values": [flt(r.get(key), 2) for r in rows]}
			]
		many = len(data["items"]) > 1
		labels, values = [], []
		for item in data["items"]:
			for w in item["warehouses"]:
				labels.append(f"{item['item_code']} · {w['warehouse']}" if many else w["warehouse"])
				values.append(flt(w["actual_qty"], 3))
		return labels[:points], [{"name": "Qty", "values": values[:points]}]

	if tool == "cash_bank_position":
		return _align(
			{co: {r["account"]: r["balance"] for r in c["accounts"]} for co, c in data["companies"].items()},
			points,
		)

	if tool == "profit_loss":
		return _align(
			{co: {g["group"]: g["amount"] for g in c["by_group"]} for co, c in data["companies"].items()},
			points,
		)

	if tool == "gst_summary":
		companies = data["companies"]
		if any("by_month" in c for c in companies.values()):
			series = {}
			for co, c in companies.items():
				prefix = f"{_abbr(co)} " if len(companies) > 1 else ""
				for kind in ("output", "input", "net"):
					series[f"{prefix}{kind.title()}"] = {m["month"]: m[kind] for m in c.get("by_month", [])}
			return _align(series, points, sort=True)
		return list(companies), [
			{"name": name, "values": [c[key] for c in companies.values()]}
			for name, key in (("Output", "output"), ("Input", "input"), ("Net payable", "net_payable"))
		]

	if tool == "payments_summary":
		group = args.get("group_by") or "account"
		label_key = {"account": "account", "mode": "mode_of_payment", "day": "posting_date"}.get(group)
		series = {}
		for co, c in data["companies"].items():
			prefix = f"{_abbr(co)} " if len(data["companies"]) > 1 else ""
			for side in ("received", "paid"):
				series[f"{prefix}{side.title()}"] = {
					str((r.get(label_key) if label_key else r.get("party_name") or r.get("party")) or "—"): r[
						"amount"
					]
					for r in c[side][f"by_{group}"]
				}
		return _align(series, points, sort=group == "day")

	if tool == "item_movement":
		if "warehouses" in data:
			rows = data["warehouses"][:points]
			return [r["warehouse"] for r in rows], [
				{"name": "Closing qty", "values": [r["closing"] for r in rows]}
			]
		# fast has a period (qty sold); not_sold is as of today (stock value, or qty without valuation).
		key = "qty_sold" if data.get("from_date") else "stock_value" if data.get("valuation_shown") else "qty"
		return _align(
			{
				co: {_label(r, "item_name"): flt(r.get(key), 2) for r in c["items"]}
				for co, c in data["companies"].items()
			},
			points,
		)

	if tool == "customer_activity":
		if data.get("customer"):
			return _align(
				{co: {m["month"]: m["total"] for m in c["by_month"]} for co, c in data["companies"].items()},
				points,
				sort=True,
			)
		return _align(
			{
				co: {_label(r, "customer_name"): r["total"] for r in c["lapsed"]}
				for co, c in data["companies"].items()
			},
			points,
		)

	raise ValueError(f"No chart for {tool}; use export_excel instead.")


def _abbr(company):
	return frappe.get_cached_value("Company", company, "abbr") or company


def _label(row, key):
	value = row.get(key) or row.get("customer") or row.get("supplier") or row.get("item_code")
	return str(value)


def _align(per_company, points, sort=False):
	"""One dataset per company over the union of labels (missing = 0)."""
	labels = []
	for values in per_company.values():
		labels += [label for label in values if label not in labels]
	if sort:
		labels.sort()
	labels = labels[-points:] if sort else labels[:points]
	datasets = [
		{"name": company, "values": [flt(values.get(label), 2) for label in labels]}
		for company, values in per_company.items()
		if values
	]
	return labels, datasets


# ── housekeeping ────────────────────────────────────────────────────────────


def delete_old_exports(days=7):
	"""Daily: export files are kept a week; the chat then says the link has expired."""
	cutoff = frappe.utils.add_days(frappe.utils.now_datetime(), -days)
	for name in frappe.get_all(
		"File",
		filters={"file_name": ["like", f"{FILE_PREFIX}%"], "creation": ["<", cutoff]},
		pluck="name",
	):
		frappe.delete_doc("File", name, ignore_permissions=True, force=True)
	frappe.db.commit()
